import pandas as pd
import tkinter as tk
import win32com.client
import os
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime
from PIL import Image, ImageTk


# ============================================= Funções =============================================

# Função para atualizar as tabelas ------------------------------------------
def atualizar_tabelas():
    # Carregar novamente os dados
    exibir_tabelas()
    messagebox.showinfo("Sucesso", "Tabelas atualizadas com sucesso!")
# ---------------------------------------------------------------------------

# Função para carregar as tabelas de Apontamentos e Refugos -----------------
def carregar_tabelas():
    try:
        # Carregar os dados das planilhas Apontamentos e Refugos
        apontamentos = pd.read_excel("Tables.xlsx", sheet_name="Apontamentos")[['Ferramenta', 'Quantidade', 'Data e Hora']]
        refugos = pd.read_excel("Tables.xlsx", sheet_name="Refugos")[['Ferramenta', 'Quantidade', 'Data e Hora']]
        return apontamentos, refugos
    except FileNotFoundError:
        messagebox.showerror("Erro", "Arquivo Tables.xlsx não encontrado.")
        return pd.DataFrame(), pd.DataFrame()
    except ValueError:
        messagebox.showerror("Erro", "Uma das planilhas Apontamentos ou Refugos não encontrada.")
        return pd.DataFrame(), pd.DataFrame()
# ---------------------------------------------------------------------------

# Função para exibir os dados nas tabelas -----------------------------------
def exibir_tabelas():
    apontamentos, refugos = carregar_tabelas()
    
    # Convertendo a coluna 'Data' para o tipo datetime, se necessário
    if 'Data e Hora' in apontamentos.columns:
        apontamentos['Data e Hora'] = pd.to_datetime(apontamentos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'apontamentos' DataFrame")

    if 'Data e Hora' in refugos.columns:
        refugos['Data e Hora'] = pd.to_datetime(refugos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'refugos' DataFrame")
    
    # Filtrando os dados para exibir apenas os que são dentro dos últimos 7 dias
    hoje = pd.to_datetime(datetime.today().date())  # Data de hoje
    sete_dias_atras = hoje - pd.Timedelta(days=7)  # Data de 7 dias atrás

    apontamentos = apontamentos[apontamentos['Data e Hora'] >= sete_dias_atras]
    refugos = refugos[refugos['Data e Hora'] >= sete_dias_atras]

    # Limpar a tabela de Apontamentos e Refugos
    for item in tree_apontamentos.get_children():
        tree_apontamentos.delete(item)
    for item in tree_refugos.get_children():
        tree_refugos.delete(item)

    # Inserir os dados de Apontamentos na tabela
    for _, row in apontamentos.iterrows():
        tree_apontamentos.insert("", "end", values=list(row))

    # Inserir os dados de Refugos na tabela
    for _, row in refugos.iterrows():
        tree_refugos.insert("", "end", values=list(row))
# ---------------------------------------------------------------------------

# Função para carregar as ferramentas da planilha ---------------------------
def carregar_ferramentas():
    try:
        sheet = pd.read_excel("Tables.xlsx", sheet_name="vw_TurnoverManagementOrders")
        ferramentas = sheet.iloc[:, 0].dropna().tolist()
        return ferramentas
    except FileNotFoundError:
        messagebox.showerror("Erro", "Arquivo Tables.xlsx não encontrado.")
        return []
    except ValueError:
        messagebox.showerror("Erro", "Planilha vw_TurnoverManagementOrders não encontrada.")
        return []
# ---------------------------------------------------------------------------

# Função para filtrar as opções do dropdown ---------------------------------
def filtrar_opcoes(event):
    texto_digitado = dropdown_var.get().lower()
    if texto_digitado == "":
        dropdown["values"] = ferramentas
    else:
        dropdown["values"] = [f for f in ferramentas if texto_digitado in f.lower()]
# ---------------------------------------------------------------------------

# Função para salvar apontamento com tipos e formatação ---------------------
def salvar_apontamento(planilha_nome):
    ferramenta = dropdown_var.get()
    quantidade = input_quantidade.get()

    if not ferramenta or not quantidade:
        messagebox.showwarning("Aviso", "Por favor, preencha todos os campos.")
        return

    try:
        quantidade = int(quantidade)
    except ValueError:
        messagebox.showerror("Erro", "A quantidade deve ser um número inteiro.")
        return

    try:
        # Definir tipos e colunas
        new_data = pd.DataFrame(
            {
                "Código": [None],  # Será preenchido com o código sequencial
                "Ferramenta": [ferramenta],
                "Quantidade": [quantidade],
                "Data e Hora": [datetime.now()],
                "Data": [datetime.now().date()],
            }
        )

        # Carregar dados existentes, se houver
        try:
            existing_data = pd.read_excel("Tables.xlsx", sheet_name=planilha_nome)
            next_code = existing_data["Código"].max() + 1 if not existing_data.empty else 1
            new_data["Código"] = next_code + new_data.index
            updated_data = pd.concat([existing_data, new_data], ignore_index=True)
        except (FileNotFoundError, ValueError):
            new_data["Código"] = 1
            updated_data = new_data

        # Salvar no Excel
        with pd.ExcelWriter("Tables.xlsx", engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            updated_data.to_excel(writer, sheet_name=planilha_nome, index=False)

        # Reaplicar a formatação como tabela
        book = load_workbook("Tables.xlsx")
        ws = book[planilha_nome]

        table_ref = f"A1:E{len(updated_data) + 1}"
        table = Table(displayName=planilha_nome, ref=table_ref)

        # Aplicar estilo de tabela
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        table.tableStyleInfo = style
        ws.add_table(table)

        # Salvar arquivo formatado
        book.save("Tables.xlsx")
        book.close()

        messagebox.showinfo("Sucesso", f"{planilha_nome} salvo com sucesso!")
        dropdown_var.set("")  # Limpar o valor selecionado do dropdown
        input_quantidade.delete(0, "end")  # Limpar o campo de quantidade

    except Exception as e:
        messagebox.showerror("Erro", f"Não foi possível salvar o apontamento: {e}")
# ---------------------------------------------------------------------------

# Função para abrir a janela de apontamento ---------------------------------
def abrir_janela_apontamento():
    global janela_apontamento, dropdown, dropdown_var, input_quantidade, ferramentas

    ferramentas = carregar_ferramentas()
    if not ferramentas:
        return

    janela_apontamento = tk.Toplevel(janela_principal)
    janela_apontamento.title("Criar Apontamento")
    janela_apontamento.geometry("460x200")  # Aumenta o tamanho da janela
    janela_apontamento.config(bg="#f0f0f0")  # Cor de fundo clara

    tk.Label(janela_apontamento, text="Ferramenta:", font=("Arial", 12), bg="#f0f0f0").grid(row=0, column=0, padx=20, pady=20, sticky="w")

    # Dropdown com funcionalidade de filtro
    dropdown_var = tk.StringVar()
    dropdown = ttk.Combobox(janela_apontamento, textvariable=dropdown_var, values=ferramentas, width=18, font=("Arial", 12))
    dropdown.grid(row=0, column=1, padx=20, pady=20)
    dropdown.bind("<KeyRelease>", filtrar_opcoes)  # Evento para detectar digitação e filtrar

    tk.Label(janela_apontamento, text="Quantidade:", font=("Arial", 12), bg="#f0f0f0").grid(row=1, column=0, padx=20, pady=10, sticky="w")
    input_quantidade = tk.Entry(janela_apontamento, font=("Arial", 12), width=20)
    input_quantidade.grid(row=1, column=1, padx=20, pady=10)

    # Botões de apontamento e refugo
    tk.Button(janela_apontamento, text="Apontar Entrada", command=lambda: salvar_apontamento("Apontamentos"), font=("Arial", 12), bg="green", fg="white", relief="flat", width=20, height=2).grid(row=2, column=0, padx=20, pady=20)
    tk.Button(janela_apontamento, text="Apontar Refugo", command=lambda: salvar_apontamento("Refugos"), font=("Arial", 12), bg="red", fg="white", relief="flat", width=20, height=2).grid(row=2, column=1, padx=20, pady=20)
# ---------------------------------------------------------------------------

# Função para criar um email no Outlook -------------------------------------
def create_outlook_email(to, subject, body, attachment_path=None):
    try:
        # Conecta ao Outlook
        outlook = win32com.client.Dispatch("Outlook.Application")
        # Cria uma mensagem de email
        mail = outlook.CreateItem(0)

        # Define os campos
        mail.To = to
        mail.Subject = subject
        mail.Body = body

        if attachment_path:
            mail.Attachments.Add(attachment_path)

        # Abre o editor de email do Outlook
        mail.Display()

    except Exception as e:
        print(f"Error: {e}")
# ---------------------------------------------------------------------------

# Função para enviar relatório ----------------------------------------------
def enviar_relatorio():
    apontamentos, refugos = carregar_tabelas()
    
    if 'Data e Hora' in apontamentos.columns:
        apontamentos['Data e Hora'] = pd.to_datetime(apontamentos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'apontamentos' DataFrame")
    
    if 'Data e Hora' in refugos.columns:
        refugos['Data e Hora'] = pd.to_datetime(refugos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'refugos' DataFrame")
    
    
    hoje = pd.to_datetime(datetime.today().date())  # Data de hoje
    # Filtrar apenas pela data, ignorando o horário
    refugos = refugos[refugos['Data e Hora'].dt.date == hoje.date()]
    
    arquivo_saida = os.path.abspath("refugos_filtrados.xlsx")
    refugos.to_excel(arquivo_saida, index=False, engine='openpyxl')
    
    
    recipient = "Destinatário"
    email_subject = "Assunto"
    email_body = "Corpo"

    create_outlook_email(recipient, email_subject, email_body, attachment_path=arquivo_saida)
# ---------------------------------------------------------------------------

# Função para atualizar os dados e exibir novamente as tabelas --------------
def refresh_data():
    
    # Limpar a tabela de Apontamentos e Refugos
    for item in tree_apontamentos.get_children():
        tree_apontamentos.delete(item)
    for item in tree_refugos.get_children():
        tree_refugos.delete(item)
    
    # Carregar novamente as tabelas de Apontamentos e Refugos
    apontamentos, refugos = carregar_tabelas()
    
    # Convertendo a coluna 'Data' para o tipo datetime, se necessário
    if 'Data e Hora' in apontamentos.columns:
        apontamentos['Data e Hora'] = pd.to_datetime(apontamentos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'apontamentos' DataFrame")

    if 'Data e Hora' in refugos.columns:
        refugos['Data e Hora'] = pd.to_datetime(refugos['Data e Hora'], errors='coerce')
    else:
        print("Coluna 'Data e Hora' não encontrada em 'refugos' DataFrame")
    
    # Filtrando os dados para exibir apenas os que são dentro dos últimos 7 dias
    hoje = pd.to_datetime(datetime.today().date())  # Data de hoje
    sete_dias_atras = hoje - pd.Timedelta(days=7)  # Data de 7 dias atrás

    apontamentos = apontamentos[apontamentos['Data e Hora'] >= sete_dias_atras]
    refugos = refugos[refugos['Data e Hora'] >= sete_dias_atras]

    
    # Inserir os dados de Apontamentos na tabela
    for _, row in apontamentos.iterrows():
        tree_apontamentos.insert("", "end", values=list(row))

    # Inserir os dados de Refugos na tabela
    for _, row in refugos.iterrows():
        tree_refugos.insert("", "end", values=list(row))
        
    messagebox.showinfo("Sucesso", "Tabelas atualizadas com sucesso!")
# ---------------------------------------------------------------------------

# Funçaõ para alterar o tamanho do cabeçalho --------------------------------
def update_header_image(event):
    # Pega o comprimeto da tela
    window_width = event.width

    # Altera o tamanho da imagem com a altura proporcional
    image_resized = image.resize((window_width, int(window_width * original_height / original_width)))

    # Converte a imagem para um formato que o Tkinter lê
    header_image_resized = ImageTk.PhotoImage(image_resized)

    # Atualiza a imagem
    header_label.config(image=header_image_resized)
    header_label.image = header_image_resized  # Mantem uma referência para a imagem
# ---------------------------------------------------------------------------

# ======================================= Estrutura das Telas =======================================

# Janela principal
janela_principal = tk.Tk()
janela_principal.title("Gestão de Apontamentos")
janela_principal.state('zoomed')
janela_principal.config(bg="#f0f0f0")  # Cor de fundo clara

image = Image.open("cabecalho.png")
original_width, original_height = image.size

image_resized = image.resize((janela_principal.winfo_screenwidth(), int(janela_principal.winfo_screenwidth() * original_height / original_width)))
header_image = ImageTk.PhotoImage(image_resized)

# Cria uma lable para mostrar a imagem como cabeçalho
header_label = tk.Label(janela_principal, image=header_image)
header_label.pack(fill="x")

# Frame para os botões na janela principal
frame_botoes = tk.Frame(janela_principal, bg="#f0f0f0")
frame_botoes.pack(pady=20)

# Botão Adicionar (ícone +)
btn_criar = tk.Button(frame_botoes, text="+", command=abrir_janela_apontamento, font=("Arial", 24), bg="#89BC6E", fg="white", relief="flat", width=4, height=1)
btn_criar.grid(row=0, column=0, padx=20)

# Botão Enviar Relatório
btn_enviar = tk.Button(frame_botoes, text="Enviar Relatório\nde refugos", command=enviar_relatorio, font=("Arial", 15), bg="#1AA4CB", fg="white", relief="flat", width=13, height=2)
btn_enviar.grid(row=0, column=1, padx=20)

# Botão de Refresh
btn_refresh = tk.Button(frame_botoes, text="Atualizar Dados", command=refresh_data, font=("Arial", 15), bg="#193D80", fg="white", relief="flat", width=13, height=2)
btn_refresh.grid(row=0, column=2, padx=20)

# Frame para as tabelas
frame_tabelas = tk.Frame(janela_principal, bg="#f0f0f0")
frame_tabelas.pack(pady=20)

# Tabela Apontamentos
frame_apontamentos = ttk.Frame(frame_tabelas)
frame_apontamentos.pack(side="left", padx=1, fill="both", expand=True)

label_apontamentos = ttk.Label(frame_apontamentos, text="Apontamentos", font=("Arial", 12, "bold"))
label_apontamentos.pack(side="top", pady=5)

tree_apontamentos = ttk.Treeview(frame_apontamentos, columns=("Ferramenta", "Quantidade", "Data e Hora"), show="headings", height=16)
tree_apontamentos.pack(side="left", padx=20)
tree_apontamentos.heading("Ferramenta", text="Ferramenta")
tree_apontamentos.heading("Quantidade", text="Quantidade")
tree_apontamentos.heading("Data e Hora", text="Data e Hora")

# Tabela Refugos
frame_refugos = ttk.Frame(frame_tabelas)
frame_refugos.pack(side="left", padx=1, fill="both", expand=True)

label_refugos = ttk.Label(frame_refugos, text="Refugos", font=("Arial", 12, "bold"))
label_refugos.pack(side="top", pady=5)

tree_refugos = ttk.Treeview(frame_refugos, columns=("Ferramenta", "Quantidade", "Data e Hora"), show="headings", height=16)
tree_refugos.pack(side="left", padx=20)
tree_refugos.heading("Ferramenta", text="Ferramenta")
tree_refugos.heading("Quantidade", text="Quantidade")
tree_refugos.heading("Data e Hora", text="Data e Hora")

# Carregar e exibir as tabelas ao iniciar
exibir_tabelas()

janela_principal.mainloop()
